"""Bulk Task Import from Excel (.xlsx).

Endpoints:
- POST /api/bulk-import/tasks/preview   – parse & validate an uploaded Excel file
- POST /api/bulk-import/tasks/execute   – import valid rows as tasks
- GET  /api/bulk-import/tasks/template  – download a blank Excel template
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime, timezone, timedelta
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func

from .. import models
from ..database import get_db
from ..deps import require_roles, get_current_user
from ..services.audit_service import log_audit

router = APIRouter(prefix="/api/bulk-import", tags=["Bulk Import"])

IST = timezone(timedelta(hours=5, minutes=30))

VALID_PRIORITIES = {"low", "medium", "high", "critical"}
PRIORITY_MAP = {"low": "Low", "medium": "Medium", "high": "High", "critical": "Critical"}

# ---------------------------------------------------------------------------
# Column-name normalisation helpers
# ---------------------------------------------------------------------------

# Map user-friendly header variants to canonical keys
_COLUMN_ALIASES: dict[str, str] = {
    "description": "description",
    "task description": "description",
    "subject": "subject",
    "task subject": "subject",
    "project": "project",
    "project name": "project",
    "module": "module",
    "main module": "module",
    "sub module": "sub_module",
    "submodule": "sub_module",
    "sub_module": "sub_module",
    "developer": "developer",
    "resource": "developer",
    "developer/resource": "developer",
    "developer / resource": "developer",
    "assigned to": "developer",
    "work type": "work_type",
    "worktype": "work_type",
    "work_type": "work_type",
    "priority": "priority",
    "status": "status",
    "start date": "start_date",
    "start_date": "start_date",
    "startdate": "start_date",
    "end date": "end_date",
    "end_date": "end_date",
    "enddate": "end_date",
    "estimated hours": "estimated_hours",
    "estimated_hours": "estimated_hours",
    "est hours": "estimated_hours",
    "est. hours": "estimated_hours",
    "hours": "estimated_hours",
    "sprint": "sprint",
    "sprint name": "sprint",
    "case ref": "case_ref",
    "case_ref": "case_ref",
    "case reference": "case_ref",
    "case#": "case_ref",
    "sfdc case#": "case_ref",
    "property": "property_client",
    "client": "property_client",
    "property/client": "property_client",
    "property_client": "property_client",
    "property / client": "property_client",
    "customer committed": "customer_committed",
    "customer_committed": "customer_committed",
    "committed": "customer_committed",
    "point of contact": "point_of_contact",
    "point_of_contact": "point_of_contact",
    "poc": "point_of_contact",
    "contact": "point_of_contact",
}

TEMPLATE_COLUMNS = [
    "Description",
    "Subject",
    "Project",
    "Module",
    "Sub Module",
    "Developer/Resource",
    "Work Type",
    "Priority",
    "Status",
    "Start Date",
    "End Date",
    "Estimated Hours",
    "Sprint",
    "Case Ref",
    "Property/Client",
    "Customer Committed",
    "Point of Contact",
]


def _normalise_headers(raw_headers: list[str]) -> dict[int, str]:
    """Return {col_index: canonical_key} mapping."""
    mapping: dict[int, str] = {}
    for idx, raw in enumerate(raw_headers):
        if raw is None:
            continue
        cleaned = str(raw).strip().lower()
        canonical = _COLUMN_ALIASES.get(cleaned)
        if canonical:
            mapping[idx] = canonical
    return mapping


# ---------------------------------------------------------------------------
# Date parsing
# ---------------------------------------------------------------------------

_DATE_FORMATS = [
    "%d/%m/%Y",
    "%Y-%m-%d",
    "%d-%b-%Y",
    "%d-%b-%y",
    "%d-%B-%Y",
    "%d-%B-%y",
    "%d.%m.%Y",
    "%m/%d/%Y",
    "%d/%m/%y",
    "%Y/%m/%d",
]


def _parse_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None  # will be flagged as unparseable


def _parse_bool(value: Any) -> bool:
    if value is None:
        return False
    s = str(value).strip().lower()
    return s in ("yes", "y", "true", "1")


def _parse_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


# ---------------------------------------------------------------------------
# Entity lookup caches (per-request)
# ---------------------------------------------------------------------------

def _build_lookup(db: Session):
    """Build case-insensitive lookup dicts for all reference entities."""
    projects = {p.name.lower(): p for p in db.query(models.Project).all()}
    developers = {d.name.lower().strip(): d for d in db.query(models.Developer).filter(models.Developer.active == True).all()}
    modules = {m.name.lower(): m for m in db.query(models.MainModule).all()}
    sub_modules = {s.name.lower(): s for s in db.query(models.SubModule).all()}
    work_types = {w.name.lower(): w for w in db.query(models.WorkType).all()}
    sprints = {s.name.lower(): s for s in db.query(models.Sprint).all()}
    return {
        "projects": projects,
        "developers": developers,
        "modules": modules,
        "sub_modules": sub_modules,
        "work_types": work_types,
        "sprints": sprints,
    }


# ---------------------------------------------------------------------------
# Row validation
# ---------------------------------------------------------------------------

def _validate_row(row_data: dict[str, Any], lookup: dict) -> tuple[dict, list[str]]:
    """Validate a single parsed row. Returns (resolved_data, errors)."""
    errors: list[str] = []
    resolved: dict[str, Any] = {}

    # --- Description (required) ---
    desc = (row_data.get("description") or "").strip() if row_data.get("description") is not None else ""
    if not desc:
        errors.append("Description is required")
    resolved["description"] = desc

    # --- Subject ---
    resolved["subject"] = (str(row_data.get("subject") or "")).strip()

    # --- Project (required) ---
    proj_name = (str(row_data.get("project") or "")).strip()
    if not proj_name:
        errors.append("Project is required")
    else:
        proj = lookup["projects"].get(proj_name.lower())
        if not proj:
            errors.append(f"Project '{proj_name}' not found")
        else:
            resolved["project_id"] = proj.id
            resolved["_project"] = proj

    # --- Developer (required) ---
    dev_name = (str(row_data.get("developer") or "")).strip()
    if not dev_name:
        errors.append("Developer/Resource is required")
    else:
        dev = lookup["developers"].get(dev_name.lower().strip())
        if not dev:
            errors.append(f"Developer '{dev_name}' not found")
        else:
            resolved["developer_id"] = dev.id

    # --- Priority (required) ---
    prio = (str(row_data.get("priority") or "")).strip()
    if not prio:
        errors.append("Priority is required (Low/Medium/High/Critical)")
    elif prio.lower() not in VALID_PRIORITIES:
        errors.append(f"Priority must be Low/Medium/High/Critical, got '{prio}'")
    else:
        resolved["priority"] = PRIORITY_MAP[prio.lower()]

    # --- Module (optional) ---
    mod_name = (str(row_data.get("module") or "")).strip()
    if mod_name:
        mod = lookup["modules"].get(mod_name.lower())
        if not mod:
            errors.append(f"Module '{mod_name}' not found")
        else:
            resolved["main_module_id"] = mod.id

    # --- Sub Module (optional) ---
    sub_name = (str(row_data.get("sub_module") or "")).strip()
    if sub_name:
        sub = lookup["sub_modules"].get(sub_name.lower())
        if not sub:
            errors.append(f"Sub Module '{sub_name}' not found")
        else:
            resolved["sub_module_id"] = sub.id

    # --- Work Type (optional) ---
    wt_name = (str(row_data.get("work_type") or "")).strip()
    if wt_name:
        wt = lookup["work_types"].get(wt_name.lower())
        if not wt:
            errors.append(f"Work Type '{wt_name}' not found")
        else:
            resolved["work_type_id"] = wt.id

    # --- Sprint (optional) ---
    sp_name = (str(row_data.get("sprint") or "")).strip()
    if sp_name:
        sp = lookup["sprints"].get(sp_name.lower())
        if not sp:
            errors.append(f"Sprint '{sp_name}' not found")
        else:
            resolved["sprint_id"] = sp.id

    # --- Status ---
    status_val = (str(row_data.get("status") or "")).strip()
    resolved["status"] = status_val if status_val else "Not Started"

    # --- Start Date ---
    raw_start = row_data.get("start_date")
    if raw_start is not None and str(raw_start).strip():
        sd = _parse_date(raw_start)
        if sd is None:
            errors.append(f"Start Date '{raw_start}' could not be parsed")
        else:
            resolved["start_date"] = sd

    # --- End Date ---
    raw_end = row_data.get("end_date")
    if raw_end is not None and str(raw_end).strip():
        ed = _parse_date(raw_end)
        if ed is None:
            errors.append(f"End Date '{raw_end}' could not be parsed")
        else:
            resolved["end_date"] = ed

    # --- Estimated Hours ---
    raw_hours = row_data.get("estimated_hours")
    if raw_hours is not None and str(raw_hours).strip():
        hours = _parse_float(raw_hours)
        if hours is None:
            errors.append(f"Estimated Hours '{raw_hours}' is not a valid number")
        else:
            resolved["estimated_hours"] = hours

    # --- Case Ref ---
    resolved["case_ref"] = (str(row_data.get("case_ref") or "")).strip()

    # --- Property/Client ---
    resolved["property_client"] = (str(row_data.get("property_client") or "")).strip()

    # --- Customer Committed ---
    resolved["customer_committed"] = _parse_bool(row_data.get("customer_committed"))

    # --- Point of Contact ---
    resolved["point_of_contact"] = (str(row_data.get("point_of_contact") or "")).strip()

    return resolved, errors


# ---------------------------------------------------------------------------
# Excel parsing helper
# ---------------------------------------------------------------------------

def _parse_excel(contents: bytes, db: Session):
    """Parse an .xlsx file and validate each row. Returns preview payload."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(500, "openpyxl is not installed on the server.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "Could not read the file. Please upload a valid .xlsx file.")

    ws = wb.active
    if ws is None:
        raise HTTPException(400, "The Excel file has no active sheet.")

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(400, "The file must have a header row and at least one data row.")

    header_row = rows[0]
    col_map = _normalise_headers([str(h) if h is not None else None for h in header_row])

    if not col_map:
        raise HTTPException(400, "No recognised columns found in the header row. Expected columns like Description, Project, Developer, Priority, etc.")

    lookup = _build_lookup(db)

    preview = []
    valid_count = 0
    invalid_count = 0

    for row_idx, row in enumerate(rows[1:], start=2):
        # Skip completely empty rows
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        row_data: dict[str, Any] = {}
        for col_idx, canonical in col_map.items():
            if col_idx < len(row):
                row_data[canonical] = row[col_idx]

        resolved, errors = _validate_row(row_data, lookup)

        # Build display data from raw values
        display_data = {}
        for canonical in ["description", "subject", "project", "module", "sub_module",
                          "developer", "work_type", "priority", "status", "start_date",
                          "end_date", "estimated_hours", "sprint", "case_ref",
                          "property_client", "customer_committed", "point_of_contact"]:
            val = row_data.get(canonical)
            display_data[canonical] = str(val).strip() if val is not None else ""

        is_valid = len(errors) == 0
        if is_valid:
            valid_count += 1
        else:
            invalid_count += 1

        preview.append({
            "row": row_idx,
            "data": display_data,
            "valid": is_valid,
            "errors": errors,
        })

    wb.close()

    return {
        "total_rows": len(preview),
        "valid_rows": valid_count,
        "invalid_rows": invalid_count,
        "preview": preview,
    }


# ---------------------------------------------------------------------------
# Task code generation (mirrors tasks.py logic)
# ---------------------------------------------------------------------------

def _generate_task_code(db: Session, sprint: models.Sprint | None) -> str:
    """Generate the next task code using MAX of the numeric suffix."""
    prefix = f"T{sprint.start_date.strftime('%y%m')}" if sprint else "T00000"
    prefix_len = len(prefix)

    from sqlalchemy import text
    # Try MSSQL syntax first, fall back to SQLite
    try:
        result = db.execute(text(
            "SELECT MAX(CAST(SUBSTRING(task_code, :plen + 1, LEN(task_code) - :plen) AS INT)) "
            "FROM PRM_tasks WHERE task_code LIKE :pattern"
        ), {"plen": prefix_len, "pattern": f"{prefix}%"})
        max_seq = result.scalar()
    except Exception:
        # SQLite fallback
        try:
            result = db.execute(text(
                "SELECT MAX(CAST(SUBSTR(task_code, :plen + 1) AS INTEGER)) "
                "FROM PRM_tasks WHERE task_code LIKE :pattern"
            ), {"plen": prefix_len, "pattern": f"{prefix}%"})
            max_seq = result.scalar()
        except Exception:
            max_seq = None

    next_seq = (max_seq or 0) + 1
    return f"{prefix}{next_seq:03d}"


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.post("/tasks/preview")
async def preview_bulk_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_roles("Admin", "Manager", "Lead")),
):
    """Upload an Excel file and get a preview with validation results."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Please upload a .xlsx file.")

    contents = await file.read()
    if len(contents) > 10 * 1024 * 1024:  # 10 MB limit
        raise HTTPException(400, "File is too large. Maximum size is 10 MB.")

    return _parse_excel(contents, db)


@router.post("/tasks/execute")
async def execute_bulk_import(
    file: UploadFile = File(...),
    rows_to_import: str = Form(default=""),  # comma-separated row numbers, or empty = all valid
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_roles("Admin", "Manager", "Lead")),
):
    """Execute the import — create tasks for the specified (or all valid) rows."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Please upload a .xlsx file.")

    contents = await file.read()
    preview_result = _parse_excel(contents, db)

    # Determine which rows to import
    selected_rows: set[int] | None = None
    if rows_to_import.strip():
        try:
            selected_rows = {int(r.strip()) for r in rows_to_import.split(",") if r.strip()}
        except ValueError:
            raise HTTPException(400, "rows_to_import must be comma-separated row numbers.")

    lookup = _build_lookup(db)
    imported = 0
    failed = 0
    error_details: list[dict] = []

    for item in preview_result["preview"]:
        row_num = item["row"]

        # Skip if specific rows requested and this isn't one
        if selected_rows is not None and row_num not in selected_rows:
            continue

        # Skip invalid rows
        if not item["valid"]:
            failed += 1
            error_details.append({"row": row_num, "errors": item["errors"]})
            continue

        try:
            # Re-resolve the data (we need the resolved IDs)
            raw_data = item["data"]
            row_data = {k: (v if v else None) for k, v in raw_data.items()}
            resolved, errors = _validate_row(row_data, lookup)

            if errors:
                failed += 1
                error_details.append({"row": row_num, "errors": errors})
                continue

            # Determine sprint for task code generation
            sprint = None
            if resolved.get("sprint_id"):
                sprint = db.get(models.Sprint, resolved["sprint_id"])

            task_code = _generate_task_code(db, sprint)

            # Remove internal helper keys
            resolved.pop("_project", None)

            task = models.Task(
                task_code=task_code,
                description=resolved.get("description", ""),
                subject=resolved.get("subject", ""),
                project_id=resolved.get("project_id"),
                main_module_id=resolved.get("main_module_id"),
                sub_module_id=resolved.get("sub_module_id"),
                developer_id=resolved.get("developer_id"),
                work_type_id=resolved.get("work_type_id"),
                sprint_id=resolved.get("sprint_id"),
                priority=resolved.get("priority", "Medium"),
                status=resolved.get("status", "Not Started"),
                start_date=resolved.get("start_date"),
                end_date=resolved.get("end_date"),
                estimated_hours=resolved.get("estimated_hours", 0),
                case_ref=resolved.get("case_ref", ""),
                property_client=resolved.get("property_client", ""),
                customer_committed=resolved.get("customer_committed", False),
                point_of_contact=resolved.get("point_of_contact", ""),
                created_at=datetime.now(IST),
            )

            # Auto-fill reporting_to from org hierarchy
            if task.developer_id and task.project_id:
                org_entry = db.query(models.OrgHierarchy).filter(
                    models.OrgHierarchy.project_id == task.project_id,
                    models.OrgHierarchy.developer_id == task.developer_id,
                ).first()
                if org_entry and org_entry.reports_to_id:
                    task.reporting_to_id = org_entry.reports_to_id

            db.add(task)
            db.flush()  # get the ID for audit logging

            log_audit(db, current_user, "CREATE", "Task", task.id, task.task_code,
                       changes={"source": {"old": None, "new": "bulk_import"}})

            imported += 1

        except Exception as e:
            failed += 1
            error_details.append({"row": row_num, "errors": [str(e)]})

    db.commit()

    return {
        "imported": imported,
        "failed": failed,
        "errors": error_details,
    }


@router.get("/tasks/template")
async def download_template(
    current_user: models.User = Depends(require_roles("Admin", "Manager", "Lead")),
):
    """Download a blank Excel template with the correct column headers and a sample row."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task Import"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, header in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Sample row
    sample = [
        "Implement user login page",    # Description
        "Login Module",                  # Subject
        "FX-POS",                        # Project
        "Authentication",                # Module
        "Login",                         # Sub Module
        "John Doe",                      # Developer/Resource
        "Development",                   # Work Type
        "High",                          # Priority
        "Not Started",                   # Status
        "01/09/2026",                    # Start Date
        "15/09/2026",                    # End Date
        8,                               # Estimated Hours
        "Sep-2026",                      # Sprint
        "CASE-001",                      # Case Ref
        "Acme Corp",                     # Property/Client
        "No",                            # Customer Committed
        "Jane Smith",                    # Point of Contact
    ]
    for col_idx, val in enumerate(sample, start=1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

    # Column widths
    widths = [35, 20, 15, 18, 18, 22, 15, 12, 15, 14, 14, 14, 14, 14, 18, 18, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Freeze header row
    ws.freeze_panes = "A2"

    # Instructions sheet
    ws2 = wb.create_sheet("Instructions")
    instructions = [
        ["Bulk Task Import – Instructions"],
        [""],
        ["1. Fill in the 'Task Import' sheet with your tasks (one row per task)."],
        ["2. Required columns: Description, Project, Developer/Resource, Priority."],
        ["3. Project names must exactly match existing projects in PRM (case-insensitive)."],
        ["4. Developer names must exactly match existing developers (case-insensitive)."],
        ["5. Priority must be one of: Low, Medium, High, Critical."],
        ["6. Dates can be in DD/MM/YYYY, YYYY-MM-DD, or DD-MMM-YYYY format."],
        ["7. If Status is left empty, it defaults to 'Not Started'."],
        ["8. Customer Committed accepts Yes/No (defaults to No)."],
        ["9. Module, Sub Module, Work Type, and Sprint are optional but must match existing values if provided."],
        [""],
        ["Column Reference:"],
        ["  Description (required) – Task description"],
        ["  Subject – Task subject"],
        ["  Project (required) – Must match an existing project name"],
        ["  Module – Must match an existing main module name"],
        ["  Sub Module – Must match an existing sub module name"],
        ["  Developer/Resource (required) – Must match an existing developer name"],
        ["  Work Type – Must match an existing work type name"],
        ["  Priority (required) – Low / Medium / High / Critical"],
        ["  Status – Default: Not Started"],
        ["  Start Date – DD/MM/YYYY, YYYY-MM-DD, or DD-MMM-YYYY"],
        ["  End Date – Same date formats as Start Date"],
        ["  Estimated Hours – Numeric value"],
        ["  Sprint – Must match an existing sprint name"],
        ["  Case Ref – Case reference string"],
        ["  Property/Client – Property or client name"],
        ["  Customer Committed – Yes / No"],
        ["  Point of Contact – Contact person name"],
    ]
    for row_idx, row in enumerate(instructions, start=1):
        cell = ws2.cell(row=row_idx, column=1, value=row[0] if row else "")
        if row_idx == 1:
            cell.font = Font(bold=True, size=14, color="4F46E5")
        elif row[0].startswith("  "):
            cell.font = Font(size=10, color="555555")
    ws2.column_dimensions["A"].width = 80

    # Write to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=PRM_Task_Import_Template.xlsx"},
    )
